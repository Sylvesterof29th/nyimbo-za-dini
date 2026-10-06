import json
import re
from pathlib import Path
from pptx import Presentation

CATEGORY = "Nyimbo za Dini"
LANGUAGE = "Swahili"


def clean_text(text):
    text = text.replace("\u00a0", " ")
    return re.sub(r"[\t ]+", " ", text).strip()


def extract_slide_text(slide):
    parts = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            parts.append(shape.text.strip())
    return "\n".join(parts)


def extract_hymn_number(filename):
    match = re.search(r"\d+", filename.stem)
    return int(match.group()) if match else None


def extract_title(text):
    lines = [clean_text(line) for line in text.splitlines() if clean_text(line)]
    if not lines:
        return "Untitled Hymn"
    return re.sub(r"[,;.!?]+$", "", lines[0]).strip()


def extract_lines(text):
    return [clean_text(line) for line in text.splitlines() if clean_text(line)]


def read_ppt_slides(file_path):
    presentation = Presentation(str(file_path))
    slide_texts = []
    for slide in presentation.slides:
        text = extract_slide_text(slide)
        if text:
            slide_texts.append(text)
    if not slide_texts:
        raise ValueError("No readable text found in presentation.")
    return slide_texts


def get_part_number(file_path):
    match = re.search(
        r"(?:part|sehemu|kipande)\s*(\d+)",
        file_path.stem,
        flags=re.IGNORECASE,
    )
    return int(match.group(1)) if match else None


def part_group_key(file_path, number):
    stem = file_path.stem.lower()
    stem = re.sub(r"\b(?:part|sehemu|kipande)\s*\d+\b", "", stem)
    stem = re.sub(r"[\s,_-]+", " ", stem).strip()
    return number, stem


def convert_single_file(file_path):
    number = extract_hymn_number(file_path)
    if number is None:
        raise ValueError(f"Could not determine hymn number from filename: {file_path.name}")

    slide_texts = read_ppt_slides(file_path)
    verses = []
    for index, slide_text in enumerate(slide_texts, start=1):
        lines = extract_lines(slide_text)
        if lines:
            verses.append({"verse": index, "lines": lines})

    if not verses:
        raise ValueError("No verses could be extracted.")

    return {
        "number": f"{number:03d}",
        "title": extract_title(slide_texts[0]),
        "category": CATEGORY,
        "language": LANGUAGE,
        "author": None,
        "composer": None,
        "verses": verses,
        "source_files": [file_path.name],
    }


def merge_part_files(file_paths):
    ordered = sorted(
        file_paths,
        key=lambda p: (get_part_number(p) or 0, p.name.casefold()),
    )
    first = convert_single_file(ordered[0])
    merged_verses = []
    sources = []
    verse_number = 1

    for file_path in ordered:
        item = convert_single_file(file_path)
        sources.append(file_path.name)
        for verse in item["verses"]:
            merged_verses.append({
                "verse": verse_number,
                "lines": verse["lines"],
            })
            verse_number += 1

    first["verses"] = merged_verses
    first["source_files"] = sources
    return first


def lyrics_fingerprint(hymn):
    text = "\n".join(
        line for verse in hymn["verses"] for line in verse["lines"]
    ).casefold()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[^\w\s]", "", text)
    return text.strip()


def validate_hymns(hymns):
    warnings = []
    seen_content = {}

    for hymn in hymns:
        if not hymn["title"].strip():
            warnings.append(f"Hymn {hymn['number']} has no title.")
        if not hymn["verses"]:
            warnings.append(f"Hymn {hymn['number']} has no verses.")
        for verse in hymn["verses"]:
            if not verse["lines"]:
                warnings.append(f"Hymn {hymn['number']} has an empty verse.")

        fingerprint = lyrics_fingerprint(hymn)
        if fingerprint:
            if fingerprint in seen_content:
                warnings.append(
                    f"Possible duplicate lyrics: '{hymn['title']}' "
                    f"({hymn['number']}) matches '{seen_content[fingerprint]}'."
                )
            else:
                seen_content[fingerprint] = f"{hymn['title']} ({hymn['number']})"

    return warnings


def convert_folder(input_folder, output_file, error_file):
    input_path = Path(input_folder)
    output_path = Path(output_file)
    error_path = Path(error_file)

    print("\n===================================")
    print("       NYIMBO ZA DINI IMPORT")
    print("===================================\n")

    if not input_path.exists():
        print(f"ERROR: Input folder does not exist: {input_path}")
        return

    ppt_files = sorted(
        list(input_path.glob("*.pptx")) + list(input_path.glob("*.ppt")),
        key=lambda p: p.name.casefold(),
    )
    if not ppt_files:
        print(f"No PowerPoint files found in: {input_path}")
        return

    errors = []
    groups = {}
    regular_files = []

    for ppt_file in ppt_files:
        number = extract_hymn_number(ppt_file)
        if number is None:
            errors.append(f"{ppt_file.name}: Could not determine hymn number.")
            continue

        if get_part_number(ppt_file) is not None:
            groups.setdefault(part_group_key(ppt_file, number), []).append(ppt_file)
        else:
            regular_files.append(ppt_file)

    hymns = []

    # Ordinary files are separate entries, even when their hymn numbers match.
    for ppt_file in regular_files:
        print(f"Processing: {ppt_file.name}")
        try:
            hymn = convert_single_file(ppt_file)
            hymns.append(hymn)
            print(f"  SUCCESS: Hymn {hymn['number']} - {hymn['title']}")
        except Exception as error:
            errors.append(f"{ppt_file.name}: {error}")
            print(f"  ERROR: {error}")

    # Explicitly numbered parts are combined in part-number order.
    for key, part_files in groups.items():
        print("Merging parts: " + ", ".join(p.name for p in part_files))
        try:
            hymn = merge_part_files(part_files)
            hymns.append(hymn)
            print(f"  SUCCESS: merged Hymn {hymn['number']} - {hymn['title']}")
        except Exception as error:
            errors.append(f"Part group {key}: {error}")
            print(f"  ERROR: {error}")

    hymns.sort(
        key=lambda h: (
            int(h["number"]),
            h["title"].casefold(),
            h["source_files"][0].casefold(),
        )
    )

    # Keep IDs as integers for compatibility with the current Flutter Hymn model.
    # Hymn number remains separate and may repeat.
    for unique_id, hymn in enumerate(hymns, start=1):
        hymn["id"] = unique_id

    warnings = validate_hymns(hymns)
    all_errors = errors + ["WARNING: " + warning for warning in warnings]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    error_path.parent.mkdir(parents=True, exist_ok=True)

    with output_path.open("w", encoding="utf-8") as file:
        json.dump(hymns, file, ensure_ascii=False, indent=2)

    with error_path.open("w", encoding="utf-8") as file:
        if all_errors:
            file.write("\n".join(all_errors) + "\n")
        else:
            file.write("No conversion errors or duplicate-content warnings.\n")

    print("\n===================================")
    print("         IMPORT COMPLETE")
    print("===================================\n")
    print(f"PowerPoint files found: {len(ppt_files)}")
    print(f"Song entries created: {len(hymns)}")
    print(f"Errors / warnings: {len(all_errors)}")
    print(f"\nJSON output: {output_path}")
    print(f"Report: {error_path}\n")
    print("Please review the report." if all_errors else "All entries passed basic validation.")


def main():
    convert_folder(
        "hymn_data/ppt",
        "hymn_data/output/hymns.json",
        "hymn_data/errors/import_errors.txt",
    )


if __name__ == "__main__":
    main()